"""
Chapter 9: Deployment and Operations Guide
Detailed academic specifications covering deployment steps, environment variables (Table 9.1),
local dual-process orchestration (start_all.bat), and production hosting architectures.
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from .helpers import (
    add_chapter_heading, add_heading2, add_heading3, add_body, add_bullet,
    set_cell_background, set_cell_margins, set_table_borders
)

def build_chapter9(doc):
    add_chapter_heading(doc, 9, "DEPLOYMENT AND OPERATIONS GUIDE")

    add_heading2(doc, "9.1 Step-by-Step Full-Stack Deployment Procedure")
    add_body(
        doc,
        "Deploying CivicSync in a clean target environment requires initializing both the Python virtual environment and "
        "the Node.js build pipeline. The sequence is structured as follows:",
        bold_prefix="Deployment Sequence: "
    )
    add_bullet(
        doc,
        "Clone the central repository from GitHub: `git clone https://github.com/saicharan0806/Nagarmitra.git` and enter the "
        "project root directory.",
        bold_prefix="Step 1 (Source Acquisition): "
    )
    add_bullet(
        doc,
        "Create an isolated Python 3.10+ virtual environment: `python -m venv .venv`. Activate the environment (`.\\.venv\\Scripts\\activate` "
        "on Windows or `source .venv/bin/activate` on Linux/macOS) and install pinned dependencies: `pip install -r backend/requirements.txt`.",
        bold_prefix="Step 2 (Python Runtime Isolation): "
    )
    add_bullet(
        doc,
        "Log into MySQL Community Server (`mysql -u root -p`), create the database `CREATE DATABASE civisync_db;`, and execute the "
        "DDL schema script: `mysql -u root -p civisync_db < database/schema.sql` to initialize tables, constraints, and seed data.",
        bold_prefix="Step 3 (Relational Schema Initialization): "
    )
    add_bullet(
        doc,
        "Navigate to `frontend/`, run `npm install` to download React, Vite, and Leaflet dependencies into `node_modules/`.",
        bold_prefix="Step 4 (Frontend Dependency Resolution): "
    )
    add_bullet(
        doc,
        "Confirm that the pre-trained weights file `yolov8n.pt` resides in `ai_service/weights/`. If absent, the system automatically "
        "downloads the checkpoint from the official Ultralytics release repository upon initial server startup.",
        bold_prefix="Step 5 (AI Checkpoint Verification): "
    )
    add_bullet(
        doc,
        "Execute `start_all.bat` on Windows or launch the backend and frontend in separate terminal multiplexers (`tmux` or separate tabs).",
        bold_prefix="Step 6 (Dual-Process Launch): "
    )

    add_heading2(doc, "9.2 Environment Configuration (.env)")
    add_body(
        doc,
        "All sensitive credentials, cryptographic salts, and operational ports are decoupled from code through environment variables:"
    )

    # Table 9.1: Environment Variable Matrix
    t_env = doc.add_table(rows=8, cols=4)
    t_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_env)

    env_headers = ["Variable Name", "Default Value", "Security Classification", "Functional Purpose"]
    env_data = [
        ("FLASK_ENV", "development", "Public / Config", "Sets execution mode (development vs. production)."),
        ("SECRET_KEY", "civicsync_prod_secret_2024", "CONFIDENTIAL", "Cryptographic key for signing session cookies and tokens."),
        ("DATABASE_URL", "mysql+pymysql://root:pass@localhost/civisync_db", "CONFIDENTIAL", "SQLAlchemy connection string with credentials."),
        ("PORT", "5000", "Public / Config", "TCP port on which the Flask WSGI server listens."),
        ("HOST", "0.0.0.0", "Public / Config", "Network interface binding (0.0.0.0 binds to all LAN interfaces)."),
        ("VITE_API_BASE_URL", "http://127.0.0.1:5000", "Public / Config", "Base URL used by the React client to dispatch API calls."),
        ("YOLO_MODEL_PATH", "ai_service/weights/yolov8n.pt", "Internal / Config", "File system path to the YOLOv8 neural network weights."),
    ]

    col_w_env = [Inches(1.8), Inches(1.8), Inches(1.6), Inches(2.0)]

    for c_idx, htext in enumerate(env_headers):
        cell = t_env.rows[0].cells[c_idx]
        cell.width = col_w_env[c_idx]
        cell.text = htext
        set_cell_background(cell, "002244")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(9.5)
        set_cell_margins(cell, top=100, bottom=100, left=70, right=70)

    for r_idx, row_values in enumerate(env_data):
        row = t_env.rows[r_idx+1]
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.width = col_w_env[c_idx]
            cell.text = val
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            p.runs[0].font.size = Pt(8.5)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            if r_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")

    add_body(doc, "Table 9.1: Environment Configuration Variables Matrix (.env File Specification).")

    add_heading2(doc, "9.3 Local Development vs. Production Hosting Architecture")
    add_body(
        doc,
        "In a local development setup, Vite runs on port 5173 and proxies `/api/*` requests to the Flask server on port 5000. "
        "In a production municipal cloud deployment, the architecture transitions to an enterprise topology:\n"
        "1. Reverse Proxy & Edge Termination: Nginx acts as the public-facing gateway, terminating SSL/TLS certificates and serving "
        "static compiled React assets from disk with gzip/brotli compression.\n"
        "2. Production WSGI Server: Flask is executed behind Gunicorn (on Linux) or Waitress (on Windows) with a multi-worker process "
        "pool (`workers = 2 * CPU_CORES + 1`) to handle concurrent tensor inferencing and file uploads without blocking.\n"
        "3. Database High Availability: MySQL runs as a managed cluster with read-replicas handling analytical queries from the "
        "administrative dashboard, while writes are directed to the primary instance."
    )
